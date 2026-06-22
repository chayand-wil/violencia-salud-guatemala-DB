#!/usr/bin/env python3
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import textwrap

FECHA_ENTREGA = "16 de abril de 2026"
TITULO_PROYECTO = "Violencia e indicadores de salud en Guatemala"
SALIDA = "Informe_Proyecto_Violencia_COMPLETO.docx"

CONSULTAS = [
    {
        "titulo": "Cantidad de homicidios por anio y departamento",
        "algebra": "γ anio, departamento; conteo -> total_homicidios (σ delito contiene 'HOMICIDIO' (hecho_delictivo ⨝ hecho ⨝ ubicacion ⨝ municipio ⨝ departamento ⨝ delito_cometido ⨝ delito))",
        "descripcion": "Agrupa los hechos con resultado de homicidio por anio y por departamento. La seleccion limita los registros al tipo de delito buscado y la fusion integra la informacion territorial y delictiva para obtener el total por cada combinacion geografica y temporal."
    },
    {
        "titulo": "Denuncias por violencia contra la mujer por municipio",
        "algebra": "γ departamento, municipio; conteo -> total_denuncias (σ tipo_hecho contiene 'MUJER' o 'VIOLENCIA CONTRA LA MUJER' o 'VCM' (denuncia ⨝ hecho ⨝ tipo_hecho ⨝ ubicacion ⨝ municipio ⨝ departamento))",
        "descripcion": "Recupera las denuncias asociadas con violencia contra la mujer y las resume por municipio. La fusion enlaza la denuncia con el hecho, el tipo de hecho y la ubicacion para medir la distribucion territorial de estos casos."
    },
    {
        "titulo": "Cinco tipos de hechos delictivos mas frecuentes en los ultimos cinco anios",
        "algebra": "cinco_primeros(γ clasificacion_delito; conteo -> total_casos (σ fecha_hecho dentro de los ultimos cinco anios (hecho_delictivo ⨝ hecho ⨝ delito_cometido ⨝ clasificacion_delito)))",
        "descripcion": "Identifica las clasificaciones delictivas con mayor frecuencia en el periodo reciente. La seleccion temporal reduce el conjunto y luego el conteo permite ordenar los resultados desde el mas recurrente hasta el menos recurrente."
    },
    {
        "titulo": "Sentencias dictadas por tipo de delito y anio",
        "algebra": "γ anio, delito; conteo -> total_sentencias (sentencia ⟕ delito)",
        "descripcion": "Resume las sentencias judiciales por anio y por delito. La fusion externa izquierda conserva todas las sentencias, incluso cuando algun registro no tiene delito asociado, para no perder trazabilidad judicial."
    },
    {
        "titulo": "Promedio de edad de victimas de violencia intrafamiliar",
        "algebra": "γ promedio(edad) -> promedio_edad_victima (σ fecha_nacimiento no es nula y fecha_hecho no es nula (caso_violencia_intrafamiliar ⨝ hecho ⨝ persona))",
        "descripcion": "Calcula la edad promedio de las victimas registradas en casos de violencia intrafamiliar. La condicion de seleccion evita registros incompletos y la agregacion resume el patron etario de las victimas."
    },
    {
        "titulo": "Distribucion de embarazos adolescentes por region",
        "algebra": "γ departamento, municipio; conteo -> total_embarazos_adolescentes (σ edad_gestante < 19 (embarazo_adolescente ⨝ hecho ⨝ ubicacion ⨝ municipio ⨝ departamento))",
        "descripcion": "Filtra los casos donde la edad de la gestante es menor de 19 anios y los agrega por departamento y municipio. El resultado permite ubicar territorialmente la incidencia de embarazos adolescentes."
    },
    {
        "titulo": "Casos de violencia infantil relacionados con trabajo infantil",
        "algebra": "γ anio, departamento; conteo -> total_casos (σ relacion con trabajo infantil = 1 (caso_violencia_ninez ⨝ hecho ⨝ ubicacion ⨝ municipio ⨝ departamento ⟕ caso_ninez_trabajo_infantil))",
        "descripcion": "Relaciona los casos de violencia en la ninez con la presencia de trabajo infantil. La consulta permite observar en que anio y en que departamento se concentran estos eventos."
    },
    {
        "titulo": "Porcentaje de denuncias por discriminacion segun etnia",
        "algebra": "γ grupo_etnico; conteo -> total_casos, porcentaje (violencia_estructural ⨝ persona ⟕ grupo_etnico)",
        "descripcion": "Mide la distribucion de denuncias o casos de discriminacion segun el grupo etnico de la victima. La fusion con el catalogo etnico permite clasificar los registros y calcular su peso relativo."
    },
    {
        "titulo": "Comparativa entre escolaridad y tipo de falta judicial",
        "algebra": "γ nivel_escolaridad, tipo_falta; conteo -> total_casos (falta ⨝ hecho ⨝ tipo_falta ⨝ involucrado_hecho ⨝ persona ⟕ nivel_escolaridad)",
        "descripcion": "Cruza el nivel de escolaridad de las personas con el tipo de falta judicial registrada. Esta relacion ayuda a describir si existe concentracion de casos en ciertos niveles educativos."
    },
    {
        "titulo": "Numero de necropsias por anio",
        "algebra": "γ anio; conteo -> total_necropsias (necropsia ⨝ hecho)",
        "descripcion": "Cuenta las necropsias registradas por anio. La consulta sirve como indicador de ocurrencia de muertes investigadas y permite ver su evolucion en el tiempo."
    },
    {
        "titulo": "Violencia estructural procesada y no procesada",
        "algebra": "γ estado_proceso; conteo -> total, porcentaje (vw_violencia_estructural_proceso)",
        "descripcion": "Clasifica los registros de violencia estructural segun hayan sido procesados o no. El resumen por estado muestra el nivel de avance del tratamiento de estos casos."
    },
    {
        "titulo": "Relacion entre tipo de empleo y ocurrencia de hechos delictivos",
        "algebra": "γ ocupacion; conteo_distinto(hecho) -> total_hechos (involucrado_hecho ⨝ persona_ocupacion ⨝ ocupacion)",
        "descripcion": "Asocia la ocupacion de las personas con la presencia de hechos delictivos en los que participaron. El conteo distinto evita duplicar un mismo hecho cuando intervienen varias personas."
    },
    {
        "titulo": "Casos de violencia contra la mujer con sentencia firme",
        "algebra": "γ anio; conteo -> total_casos (σ fallo firme y relacion con violencia contra la mujer (sentencia ⟕ tipo_fallo ⟕ tipo_sentencia))",
        "descripcion": "Filtra las sentencias que corresponden a resoluciones firmes vinculadas con violencia contra la mujer. Esto permite medir la respuesta judicial sobre esta problematica."
    },
    {
        "titulo": "Idiomas hablados por victimas de discriminacion",
        "algebra": "γ idioma; conteo -> total_victimas (violencia_estructural ⨝ idioma_persona ⨝ idioma_lengua)",
        "descripcion": "Identifica los idiomas reportados por las victimas de discriminacion y los resume por frecuencia. La informacion ayuda a reconocer la dimension cultural del problema."
    },
    {
        "titulo": "Numero de personas por hogar en casos de violencia intrafamiliar",
        "algebra": "γ caso_violencia_intrafamiliar, hogar; conteo_distinto(persona) -> total_personas_hogar (caso_vif_hogar ⟕ persona_hogar)",
        "descripcion": "Estima cuantas personas habitan en los hogares vinculados con casos de violencia intrafamiliar. La media del conteo por hogar ayuda a describir el entorno de convivencia."
    },
    {
        "titulo": "Violencia infantil escolarizada y no escolarizada",
        "algebra": "γ condicion_escolar; conteo -> total_casos, porcentaje (caso_violencia_ninez ⟕ estado_escolarizacion)",
        "descripcion": "Agrupa los casos de violencia infantil segun la condicion de escolarizacion. El resultado permite comparar la presencia de casos entre poblacion escolarizada y no escolarizada."
    },
    {
        "titulo": "Casos de trabajo infantil por sector economico",
        "algebra": "γ sector_economico; conteo -> total_casos (caso_ninez_trabajo_infantil ⟕ sector_economico)",
        "descripcion": "Clasifica los casos de trabajo infantil por sector economico. Este cruce ayuda a detectar en que ramas de actividad se concentra con mayor frecuencia la explotacion infantil."
    },
    {
        "titulo": "Relacion entre edad y tipo de violencia sufrida",
        "algebra": "union de tres subconjuntos: γ tipo_violencia, subtipo; promedio(edad), minimo(edad), maximo(edad), conteo ( (caso_violencia_intrafamiliar ⨝ hecho ⨝ persona) ∪ (caso_violencia_ninez ⨝ hecho ⨝ persona) ∪ (violencia_estructural ⨝ hecho ⨝ persona) )",
        "descripcion": "Combina tres tipos de violencia para estudiar la edad de las victimas y su relacion con la forma de agresion sufrida. La union permite comparar patrones etarios entre violencia intrafamiliar, violencia contra la ninez y violencia estructural."
    },
    {
        "titulo": "Desnutricion aguda por departamento, municipio y anio",
        "algebra": "γ anio, departamento, municipio; sumatoria(casos) -> total_casos (vw_registro_salud_territorial con indicador de desnutricion aguda)",
        "descripcion": "Resume la desnutricion aguda por territorio y por anio. La agregacion territorial permite observar donde se concentran los mayores valores del indicador sanitario."
    },
    {
        "titulo": "Retardo en el desarrollo por grupo etario, sexo y region",
        "algebra": "γ anio, departamento, grupo_etario, sexo; sumatoria(casos) -> total_casos (vw_registro_salud_territorial con indicador de retardo en el desarrollo)",
        "descripcion": "Cruza el retardo en el desarrollo con el grupo etario, el sexo y la region. Esto facilita comparar como se distribuye el problema en la poblacion infantil y adolescente."
    },
    {
        "titulo": "Incidencia de enfermedades cronicas por codigo CIE-10",
        "algebra": "γ anio, codigo_cie10, condicion; sumatoria(casos), tasa por cien mil (registro_salud ⨝ condicion_salud ⟕ denominador_poblacional)",
        "descripcion": "Relaciona los casos de enfermedades cronicas con su clasificacion CIE-10 y con la poblacion de referencia. La tasa ajustada permite comparar territorios con distinta poblacion."
    },
    {
        "titulo": "Evolucion de dengue y dengue grave entre 2012 y 2024",
        "algebra": "γ anio, indicador_salud; sumatoria(casos) -> total_casos (vw_registro_salud_territorial con indicador de dengue)",
        "descripcion": "Observa la evolucion temporal de dengue y dengue grave. El resumen por anio permite reconocer alzas, caidas y posibles periodos de mayor incidencia."
    },
    {
        "titulo": "Malaria por grupo etario y sexo en el ambito municipal",
        "algebra": "γ anio, departamento, municipio, grupo_etario, sexo; sumatoria(casos) -> total_casos (vw_registro_salud_territorial con indicador de malaria)",
        "descripcion": "Desglosa los casos de malaria por territorio, edad y sexo. El cruce permite analizar con mayor precision la distribucion de esta enfermedad transmitida por vectores."
    },
    {
        "titulo": "Desnutricion infantil frente a violencia intrafamiliar por departamento",
        "algebra": "fusion externa completa entre dos agregados: desnutricion por anio y departamento, y violencia intrafamiliar por anio y departamento",
        "descripcion": "Compara dos fenomenos que pueden relacionarse en el territorio: desnutricion y violencia intrafamiliar. La fusion externa completa conserva los departamentos que aparecen en uno u otro conjunto, para no perder informacion."
    },
    {
        "titulo": "Top cinco municipios con mayor suma de Chagas, Zika y Chikungunya",
        "algebra": "cinco_primeros(γ departamento, municipio; sumatoria(casos) -> total_casos (vw_registro_salud_territorial con indicadores de Chagas, Zika y Chikungunya))",
        "descripcion": "Suma tres enfermedades transmitidas por vectores para detectar los municipios con mayor carga conjunta. El resultado prioriza los territorios con mayor necesidad de intervencion sanitaria."
    },
    {
        "titulo": "Vectores y urbanizacion por departamento",
        "algebra": "γ anio, departamento, area_geografica; sumatoria(casos) -> total_casos (registro_salud ⨝ ubicacion ⨝ municipio ⨝ departamento ⟕ area_geografica)",
        "descripcion": "Cruza las enfermedades transmitidas por vectores con el tipo de area geografica. Esto ayuda a observar si existen diferencias entre zonas urbanas y rurales en la distribucion de los casos."
    },
    {
        "titulo": "Tasa de morbilidad materna infantil por anio y departamento",
        "algebra": "γ anio, departamento; sumatoria(casos), tasa por cien mil (registro_salud ⨝ ubicacion ⨝ municipio ⨝ departamento ⟕ denominador_poblacional)",
        "descripcion": "Resume la morbilidad materna infantil por territorio y anio. La tasa ajustada facilita la comparacion entre departamentos con distinto tamano poblacional."
    },
    {
        "titulo": "Correlacion entre desnutricion y hechos delictivos por municipio",
        "algebra": "union de dos series por anio, departamento y municipio: desnutricion y hechos delictivos; luego emparejamiento por claves territoriales",
        "descripcion": "Construye dos series territoriales comparables, una de salud y otra de violencia, para estudiar su relacion en el mismo espacio y periodo. La salida deja la base lista para analisis de asociacion posterior."
    },
    {
        "titulo": "Violencia estructural en area urbana y rural",
        "algebra": "γ anio, area_geografica; conteo -> total_casos (violencia_estructural ⨝ hecho ⨝ ubicacion ⟕ area_geografica)",
        "descripcion": "Clasifica la violencia estructural segun el area geografica donde ocurre. Esta consulta ayuda a reconocer si el fenomeno se concentra en zonas urbanas o rurales."
    },
    {
        "titulo": "Frecuencia de exhumaciones por tipo de delito relacionado",
        "algebra": "γ delito; conteo -> total_exhumaciones (exhumacion ⨝ hecho ⟕ hecho_delictivo ⟕ delito_cometido ⟕ delito)",
        "descripcion": "Relaciona las exhumaciones con el tipo de delito asociado al hecho. El cruce permite observar si existe una orientacion del trabajo forense hacia ciertos delitos."
    }
]


def limpiar_formato(parrafo):
    for run in parrafo.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(11)


def agregar_parrafo(doc, texto, negrita=False, cursiva=False, tamano=11, alineacion=None):
    p = doc.add_paragraph()
    if alineacion is not None:
        p.alignment = alineacion
    r = p.add_run(texto)
    r.font.name = "Calibri"
    r.font.size = Pt(tamano)
    r.font.bold = negrita
    r.font.italic = cursiva
    return p


def agregar_titulo(doc, texto, nivel=1):
    p = doc.add_heading(texto, level=nivel)
    for run in p.runs:
        run.font.name = "Calibri"
    return p


def agregar_bala(doc, texto):
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(texto)
    r.font.name = "Calibri"
    r.font.size = Pt(11)
    return p


def agregar_bloque_algebra(doc, texto):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.75)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(textwrap.fill(texto, width=92))
    r.font.name = "Courier New"
    r.font.size = Pt(9)

    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "F2F2F2")
    p._element.get_or_add_pPr().append(shading)
    return p


def agregar_espacio(doc, cantidad=1):
    for _ in range(cantidad):
        doc.add_paragraph("")


def configurar_doc(doc):
    for section in doc.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(2.54)
        section.right_margin = Cm(2.54)

    estilos = doc.styles
    estilos["Normal"].font.name = "Calibri"
    estilos["Normal"].font.size = Pt(11)

    for nombre, tamano, color in [("Heading 1", 16, RGBColor(0, 51, 102)), ("Heading 2", 13, RGBColor(0, 72, 128)), ("Heading 3", 11, RGBColor(0, 72, 128))]:
        estilo = estilos[nombre]
        estilo.font.name = "Calibri"
        estilo.font.size = Pt(tamano)
        estilo.font.bold = True
        estilo.font.color.rgb = color


def portada(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("UNIVERSIDAD DE SAN CARLOS DE GUATEMALA")
    r.font.name = "Calibri"
    r.font.size = Pt(14)
    r.font.bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("DIVISIÓN DE CIENCIAS DE LA INGENIERÍA").font.size = Pt(12)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("ÁREA PROFESIONAL").font.size = Pt(12)

    agregar_espacio(doc, 2)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("BASES DE DATOS 1\nProyecto 26\n\n")
    r.font.name = "Calibri"
    r.font.size = Pt(16)
    r.font.bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(TITULO_PROYECTO)
    r.font.name = "Calibri"
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 51, 102)

    agregar_espacio(doc, 2)

    agregar_parrafo(doc, "Docente: Ing. Pedro Domingo", alineacion=WD_ALIGN_PARAGRAPH.CENTER)
    agregar_parrafo(doc, "Estudiante: Wilson Jonatan Chay Santizo", alineacion=WD_ALIGN_PARAGRAPH.CENTER)
    agregar_parrafo(doc, f"Fecha de entrega: {FECHA_ENTREGA}", alineacion=WD_ALIGN_PARAGRAPH.CENTER)

    doc.add_page_break()


def indice_manual(doc):
    agregar_titulo(doc, "Indice", 1)
    agregar_bala(doc, "Introducción")
    agregar_bala(doc, "Justificación")
    agregar_bala(doc, "Arquitectura propuesta")
    agregar_bala(doc, "Modelo de datos y análisis")
    agregar_bala(doc, "Proceso ETL")
    agregar_bala(doc, "Consultas analíticas y álgebra relacional")
    agregar_bala(doc, "Conclusiones")
    doc.add_page_break()


def seccion_introduccion(doc):
    agregar_titulo(doc, "Introducción", 1)
    texto = (
        "El proyecto integra información de violencia y salud pública para Guatemala con el fin de "
        "construir una base de datos normalizada que permita consultas analíticas a nivel departamental "
        "y municipal. La idea central no es solo almacenar registros, sino relacionar eventos, personas, "
        "ubicaciones y clasificaciones de salud y violencia dentro de un mismo modelo."
    )
    agregar_parrafo(doc, texto)
    texto = (
        "La propuesta parte de datos heterogéneos y, en varios casos, incompletos. Por esa razón se optó "
        "por un diseño que preserve la trazabilidad de cada evento aunque falten algunos atributos. "
        "La estructura relacional facilita el cruce entre fuentes y permite comparar fenómenos sociales "
        "que suelen estudiarse por separado."
    )
    agregar_parrafo(doc, texto)

    agregar_titulo(doc, "Objetivo general", 2)
    agregar_parrafo(doc, "Diseñar e implementar una base de datos normalizada que integre indicadores de violencia e indicadores de salud pública en Guatemala para facilitar análisis relacionales entre ambas dimensiones.")

    agregar_titulo(doc, "Alcance del proyecto", 2)
    for item in [
        "Violencia delictiva, intrafamiliar, estructural, contra la mujer y contra la niñez.",
        "Indicadores de desnutrición, enfermedades crónicas, enfermedades transmitidas por vectores y morbilidad materna infantil.",
        "Cobertura territorial nacional con referencia departamental y municipal.",
        "Integración de fuentes oficiales y conservación del dato original de cada registro."
    ]:
        agregar_bala(doc, item)


def seccion_justificacion(doc):
    agregar_titulo(doc, "Justificación", 1)
    agregar_parrafo(doc, "La violencia y la salud pública no se comportan como problemas aislados. En Guatemala comparten determinantes sociales como pobreza, desigualdad territorial, exclusión educativa, precariedad laboral y falta de acceso a servicios básicos. Un modelo de datos unificado permite observar esas relaciones de manera consistente.")
    agregar_parrafo(doc, "La utilidad principal del sistema está en que concentra información de diferentes instituciones y la organiza bajo claves comunes de territorio, tiempo y persona. Con ello se puede estudiar si ciertos municipios presentan al mismo tiempo mayor violencia y peores indicadores de salud, o si existen patrones que se repiten en determinados grupos poblacionales.")

    agregar_titulo(doc, "Sentido analítico", 2)
    for item in [
        "Permite identificar coincidencias territoriales entre violencia y salud.",
        "Facilita el estudio de grupos etarios, sexo, escolaridad, ocupación e idioma.",
        "Ayuda a justificar decisiones públicas a partir de evidencia integrada.",
        "Reduce la fragmentación de información al normalizar catálogos y eventos."
    ]:
        agregar_bala(doc, item)


def seccion_arquitectura(doc):
    agregar_titulo(doc, "Arquitectura propuesta", 1)
    agregar_titulo(doc, "Tecnologías de trabajo", 2)
    for item in [
        "Firebird como gestor de base de datos.",
        "Scripts SQL para la creación y carga de tablas.",
        "Procesos de apoyo para extracción, limpieza y homologación.",
        "Hojas de cálculo y archivos fuente de instituciones públicas."
    ]:
        agregar_bala(doc, item)

    agregar_titulo(doc, "Diseño general", 2)
    agregar_parrafo(doc, "La arquitectura se organiza en tres capas: fuentes de datos, proceso de preparación y base relacional. Primero se identifican los archivos de violencia y salud; luego se depuran, homologan y transforman; finalmente se cargan en un esquema relacional preparado para análisis.")

    agregar_titulo(doc, "Modelo lógico", 2)
    agregar_parrafo(doc, "El modelo lógico parte de las entidades persona, hecho, ubicación, delito y catálogos territoriales. A partir de ellas se conectan tablas transaccionales como denuncia, sentencia, hecho delictivo, casos de violencia intrafamiliar, violencia contra la niñez, violencia estructural y registros sanitarios. La relación entre tablas busca que cada evento quede ubicado en el tiempo, en el territorio y en su clasificación correspondiente.")


def seccion_modelo(doc):
    agregar_titulo(doc, "Modelo de datos y análisis", 1)
    agregar_titulo(doc, "Normalización", 2)
    agregar_parrafo(doc, "La base fue pensada para mantenerse en tercera forma normal. Las tablas de catálogo almacenan descripciones estables; las tablas transaccionales guardan hechos concretos; y las tablas de relación permiten registrar distintos papeles de una misma persona dentro de un evento. Con ello se evita repetir información y se reduce la ambigüedad.")

    agregar_titulo(doc, "Criterio para datos incompletos", 2)
    for item in [
        "Si falta un dato, se conserva el evento y se completa solo lo indispensable para no perder trazabilidad.",
        "Si no existe un código de catálogo, se usa el valor de desconocido.",
        "Si no viene el municipio, se toma un municipio cabecera del departamento.",
        "Si una fecha debe ser inventada, se respeta la cronología del registro fuente.",
        "Si un dato personal falta por completo, puede completarse de forma sintética solo cuando la estructura del proyecto lo requiera."
    ]:
        agregar_bala(doc, item)

    agregar_titulo(doc, "Lectura del modelo", 2)
    agregar_parrafo(doc, "El modelo permite relacionar un mismo hecho con personas de diferente rol, con delitos o tipos de hecho, con la ubicación geográfica y con variables sociales o sanitarias. Esa capacidad de cruce es la base de los indicadores que se construyen después en las consultas analíticas.")


def seccion_etl(doc):
    agregar_titulo(doc, "Proceso ETL", 1)
    agregar_titulo(doc, "Fuentes de información", 2)
    for item in [
        "Denuncias del Ministerio Público por violencia contra la mujer.",
        "Hechos delictivos contra mujeres.",
        "Medidas de seguridad.",
        "Sentencias del Ministerio Público por violencia contra la mujer.",
        "Sentencias del Organismo Judicial por violencia contra la mujer.",
        "Atenciones brindadas por el Instituto de la Víctima.",
        "Casos de discriminación.",
        "Detenidos de la Policía Nacional Civil.",
        "Víctimas de la Policía Nacional Civil.",
        "Agraviados.",
        "Sindicados.",
        "Necropsias.",
        "Exhumaciones.",
        "Médicos del INACIF.",
        "Sentenciados del Organismo Judicial.",
        "Quejas del Ministerio de Educación.",
        "Violencia intrafamiliar 2023.",
        "Violencia intrafamiliar 2024.",
        "Desnutrición.",
        "Enfermedades crónicas.",
        "Enfermedades transmitidas por vectores.",
        "Morbilidad materna infantil."
    ]:
        agregar_bala(doc, item)

    agregar_titulo(doc, "Tratamiento general", 2)
    agregar_parrafo(doc, "La extracción se orientó a conservar el dato original de cada fuente. Durante la transformación se homologaron nombres, se unificaron catálogos y se resolvieron diferencias de clasificación entre instituciones. En la carga se respetó el orden lógico del modelo para no romper dependencias entre catálogos, entidades base y tablas transaccionales.")


def seccion_consultas(doc):
    agregar_titulo(doc, "Consultas analíticas y álgebra relacional", 1)
    agregar_parrafo(doc, "Las consultas siguientes se redactaron a partir del archivo de consultas analíticas del proyecto. En cada caso se muestra la idea relacional principal y una descripción breve del cruce de datos.")

    for numero, consulta in enumerate(CONSULTAS, start=1):
        agregar_titulo(doc, f"Consulta {numero}. {consulta['titulo']}", 2)
        p = doc.add_paragraph()
        r1 = p.add_run("Álgebra relacional")
        r1.bold = True
        r1.font.name = "Calibri"
        r1.font.size = Pt(11)

        agregar_bloque_algebra(doc, consulta["algebra"])

        p = doc.add_paragraph()
        r1 = p.add_run("Descripción")
        r1.bold = True
        r1.font.name = "Calibri"
        r1.font.size = Pt(11)
        agregar_parrafo(doc, consulta["descripcion"])

        agregar_espacio(doc, 1)


def seccion_conclusiones(doc):
    agregar_titulo(doc, "Conclusiones", 1)
    agregar_parrafo(doc, "El proyecto permite ver que la violencia y la salud pública forman parte de una misma realidad social. El valor principal del modelo no está solo en almacenar registros, sino en relacionarlos para que cada hecho pueda leerse en su contexto territorial, temporal y social. Esa integración hace posible estudiar de forma conjunta fenómenos que antes estaban dispersos en archivos distintos.")
    agregar_parrafo(doc, "Las consultas construidas muestran el potencial del esquema: algunas comparan violencia con territorio, otras cruzan características de las personas involucradas y otras relacionan violencia con salud. En conjunto, estas consultas permiten observar patrones que ayudan a identificar zonas con mayor vulnerabilidad y a entender mejor cómo se conectan los distintos indicadores del proyecto.")
    agregar_parrafo(doc, "En síntesis, la base de datos propuesta ofrece una estructura consistente para consolidar información oficial, normalizar datos heterogéneos y producir análisis útiles para la interpretación de la violencia y de la salud pública en Guatemala. El cruce entre ambas dimensiones fortalece la lectura del problema social y amplía las posibilidades de análisis posterior.")


def main():
    doc = Document()
    configurar_doc(doc)
    portada(doc)
    indice_manual(doc)
    seccion_introduccion(doc)
    seccion_justificacion(doc)
    seccion_arquitectura(doc)
    seccion_modelo(doc)
    seccion_etl(doc)
    seccion_consultas(doc)
    seccion_conclusiones(doc)

    import os
    os.makedirs(os.path.dirname(SALIDA) or ".", exist_ok=True)
    doc.save(SALIDA)
    print(f"Documento generado: {SALIDA}")


if __name__ == "__main__":
    main()
